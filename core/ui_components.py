"""
共享 UI 组件 · 专业金融主题
============================
提供各 Streamlit 页面复用的 UI 组件，统一视觉风格。
"""

import streamlit as st
import pandas as pd
from typing import Any, Dict, List, Optional

from core.ui_theme import apply_plotly_theme, page_header_colors, surface_colors, status_badge_palette

# Plotly 主题由 ui_theme 按浅色/深色切换
apply_plotly_theme()


# ============================================================
# 分组 accent 色板（4 档收敛）
# ============================================================

GROUP_ACCENTS = {
    "market": "#2563EB",     # 市场研判
    "screen": "#2563EB",     # 基金筛选
    "agent": "#4F46E5",      # Agent 分析
    "monitor": "#D97706",    # 监控预警
    "portfolio": "#0F766E",  # 组合管理
    "tools": "#0F766E",      # 常用工具
    "quant": "#4F46E5",      # 量化分析
    "account": "#0F766E",    # 账户
}

# 各页历史 accent_color → 分组键，由 render_page_header 自动归一
_LEGACY_ACCENT_TO_GROUP = {
    "#3B82F6": "screen",
    "#0EA5E9": "market",
    "#7C3AED": "agent",
    "#F59E0B": "monitor",
    "#10B981": "portfolio",
    "#0891B2": "tools",
    "#6366F1": "quant",
    "#8B5CF6": "quant",
}


def _resolve_accent_color(group: Optional[str] = None, accent_color: Optional[str] = None) -> str:
    if group:
        return GROUP_ACCENTS.get(group, GROUP_ACCENTS["screen"])
    if accent_color:
        mapped = _LEGACY_ACCENT_TO_GROUP.get(accent_color.upper())
        if mapped:
            return GROUP_ACCENTS[mapped]
    return GROUP_ACCENTS["screen"]


# ============================================================
# 页面标题
# ============================================================

def render_page_header(title: str, icon: str, description: str, help_text: Optional[str] = None,
                       accent_color: Optional[str] = None, group: Optional[str] = None):
    """统一的页面标题区域 — 左侧色块 + 标题 + 描述。

    group 未传时按调用方文件名从 nav_catalog 自动映射。
    """
    # 页面级认证守卫（SQLite 自动登录本地用户；PG 多用户要求登录）
    # 注意：require_auth 未认证时会 show_login_page()+st.stop()，不可用 try/except 吞掉。
    from core.middleware import require_auth
    require_auth()
    apply_plotly_theme()

    if not group:
        try:
            import inspect
            from .nav_catalog import group_for_page
            caller = inspect.stack()[1].filename
            group = group_for_page(caller)
        except Exception:
            group = None

    color = _resolve_accent_color(group=group, accent_color=accent_color)
    hc = page_header_colors()
    safe_desc = description.replace("<", "&lt;").replace(">", "&gt;")
    st.markdown(f"""
    <div style="display:flex;align-items:flex-start;gap:1.25rem;
                background:{hc['card_bg']};border:1px solid {hc['border']};border-radius:12px;
                padding:1.5rem 1.75rem;margin-bottom:1.5rem;
                border-left:4px solid {color};box-shadow:{hc['shadow']};">
        <div style="font-size:2rem;flex-shrink:0;width:48px;height:48px;display:flex;
                    align-items:center;justify-content:center;
                    background:{color}0D;border-radius:10px;">{icon}</div>
        <div style="flex:1;">
            <div style="font-size:1.2rem;font-weight:700;color:{hc['title']};margin-bottom:0.3rem;">{title}</div>
            <div style="font-size:0.85rem;color:{hc['desc']};line-height:1.6;">{safe_desc}</div>
        </div>
    </div>
    """, unsafe_allow_html=True)

    if help_text:
        with st.expander("使用说明", expanded=False):
            st.markdown(help_text)
        st.markdown("")


# ============================================================
# 状态徽章
# ============================================================

STATUS_COLORS = {
    "success": ("#059669", "#ECFDF5", "✅ 正常"),
    "warning": ("#D97706", "#FFFBEB", "⚠️ 关注"),
    "danger":  ("#DC2626", "#FEF2F2", "❌ 风险"),
    "info":    ("#2563EB", "#EFF6FF", "ℹ️ 信息"),
    "neutral": ("#64748B", "#F8FAFC", "—"),
}

def render_status_badge(status: str) -> str:
    """渲染状态徽章 HTML。"""
    palette = status_badge_palette()
    color, bg, label = palette.get(status, palette["neutral"])
    return (
        f'<span style="display:inline-block;padding:3px 10px;border-radius:20px;'
        f'background:{bg};color:{color};font-size:0.8rem;font-weight:500;">{label}</span>'
    )


def render_risk_badge(level: str) -> str:
    """渲染风险等级标签。"""
    mapping = {
        "高":   ("#DC2626", "#FEF2F2", "高风险"),
        "中高": ("#EA580C", "#FFF7ED", "中高风险"),
        "中":   ("#D97706", "#FFFBEB", "中风险"),
        "中低": ("#059669", "#ECFDF5", "中低风险"),
        "低":   ("#10B981", "#ECFDF5", "低风险"),
    }
    color, bg, text = mapping.get(level, ("#64748B", "#F8FAFC", level))
    return (
        f'<span style="display:inline-block;padding:3px 10px;border-radius:20px;'
        f'background:{bg};color:{color};font-size:0.8rem;font-weight:600;">{text}</span>'
    )


# ============================================================
# 指标/数据展示
# ============================================================

def show_dataframe(df: pd.DataFrame, title: Optional[str] = None, use_container_width: bool = True):
    """安全展示 DataFrame。"""
    if df is None or df.empty:
        st.info("暂无数据")
        return
    if title:
        st.markdown(f"**{title}**（{len(df)} 条）")
    st.dataframe(df, use_container_width=use_container_width, hide_index=True,
                 height=min(400, 35 * len(df) + 38))


def render_metrics_row(metrics: List[Dict[str, Any]], columns: int = 4):
    """渲染一行指标卡片。"""
    cols = st.columns(columns)
    for i, m in enumerate(metrics):
        with cols[i % columns]:
            st.metric(label=m.get("label", ""), value=m.get("value", "—"),
                      delta=m.get("delta", None), help=m.get("help", None))


def render_kv_cards(items: List[tuple], columns: int = 4):
    """渲染「标签: 值」文本信息卡（替代 st.metric 显示纯文本，避免截断与样式错位）。

    Args:
        items: [(label, value), ...]
        columns: 每行列数
    """
    cols = st.columns(columns)
    sc = surface_colors()
    for i, (label, value) in enumerate(items):
        with cols[i % columns]:
            val = "—" if value in (None, "", "N/A") else str(value)
            st.markdown(
                f"""<div style="background:{sc['card_bg']};border:1px solid {sc['border']};border-radius:8px;
                        padding:0.6rem 0.9rem;box-shadow:0 1px 2px rgba(0,0,0,0.04);margin-bottom:0.5rem;">
                    <div style="font-size:0.72rem;color:{sc['label']};margin-bottom:3px;">{label}</div>
                    <div style="font-size:0.98rem;font-weight:600;color:{sc['value']};
                                overflow:hidden;text-overflow:ellipsis;white-space:nowrap;"
                         title="{val}">{val}</div>
                </div>""",
                unsafe_allow_html=True,
            )


# ============================================================
# 金融专业风：分区标题 / KPI / 报价条
# ============================================================

def section_header(title: str, subtitle: str = "") -> None:
    """紧凑分区标题（数据优先）。"""
    sub = f'<div class="fin-sec-sub">{subtitle}</div>' if subtitle else ""
    st.markdown(
        f'<div class="fin-sec"><div class="fin-sec-title">{title}</div>{sub}</div>',
        unsafe_allow_html=True,
    )


def fin_pill(text: str, tone: str = "") -> str:
    cls = f"fin-pill {tone}".strip()
    return f'<span class="{cls}">{text}</span>'


def fin_callout(text: str, tone: str = "info") -> None:
    cls = "fin-callout" if tone == "info" else f"fin-callout {tone}"
    st.markdown(f'<div class="{cls}">{text}</div>', unsafe_allow_html=True)


def kpi_strip(items: List[Dict[str, Any]]) -> None:
    """
    紧凑 KPI 条。
    items: [{label, value, delta?, tone?}] tone=up|down|muted
    """
    if not items:
        return
    cells = []
    for it in items:
        delta = it.get("delta")
        tone = it.get("tone") or ""
        if not tone and isinstance(delta, str):
            if delta.startswith("+"):
                tone = "up"
            elif delta.startswith("-"):
                tone = "down"
        d_html = f'<div class="d {tone}">{delta}</div>' if delta not in (None, "") else ""
        cells.append(
            f'<div class="fin-kpi"><div class="l">{it.get("label","")}</div>'
            f'<div class="v">{it.get("value","—")}</div>{d_html}</div>'
        )
    st.markdown(f'<div class="fin-kpi-row">{"".join(cells)}</div>', unsafe_allow_html=True)


def quote_strip(
    quotes: Dict[str, dict],
    *,
    price_key: str = "最新价",
    chg_key: str = "涨跌幅",
    price_fmt: str = "{:.2f}",
) -> None:
    """指数/风格报价条。quotes: {名称: {最新价, 涨跌幅}}"""
    if not quotes:
        st.caption("行情暂不可用")
        return
    cells = []
    for name, d in quotes.items():
        if not isinstance(d, dict):
            continue
        chg = float(d.get(chg_key) or 0)
        price = float(d.get(price_key) or 0)
        tone = "up" if chg >= 0 else "down"
        cls = f"c fin-{tone}"
        cells.append(
            f'<div class="fin-quote {tone}">'
            f'<div class="n">{name}</div>'
            f'<div class="p">{price_fmt.format(price)}</div>'
            f'<div class="{cls}">{chg:+.2f}%</div></div>'
        )
    st.markdown(f'<div class="fin-quote-row">{"".join(cells)}</div>', unsafe_allow_html=True)


def compact_fig(fig, height: int = 300, title: Optional[str] = None):
    """统一图表留白与字号。"""
    if title:
        fig.update_layout(title=dict(text=title, font=dict(size=13)))
    has_title = bool(title) or bool(getattr(fig.layout, "title", None) and fig.layout.title.text)
    fig.update_layout(
        height=height,
        margin=dict(l=8, r=8, t=40 if has_title else 16, b=8),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, font=dict(size=11)),
    )
    apply_plotly_theme()
    st.plotly_chart(fig, use_container_width=True)


def show_sortable_df(df: pd.DataFrame, title: Optional[str] = None,
                     height: Optional[int] = None,
                     column_config: Optional[dict] = None):
    """展示表格，并自动把「百分比/小数」字符串列转为数值，保证表头排序正确。

    - "12.34%" / "-3.5%" 之类的列 → 数值 + NumberColumn(%.2f%%)
    - "1.23" 之类含小数点的列 → 数值 + NumberColumn(%.2f)
    - 基金代码等纯整数/带前导零的字符串不会被转换。
    展示时仍保留原有的百分号与小数位。
    """
    import re as _re
    if df is None or df.empty:
        st.info("暂无数据")
        return

    disp = df.copy()
    cfg = dict(column_config or {})
    _id_kw = ("代码", "编号", "名称", "简称", "经理", "公司", "日期", "类型", "code")

    for col in disp.columns:
        if col in cfg:
            continue
        if any(kw in str(col) for kw in _id_kw):
            continue
        s = disp[col]
        if s.dtype != object:
            continue
        vals = s.astype(str).str.strip()
        pct_mask = vals.str.fullmatch(r"[+-]?\d+(?:\.\d+)?%").fillna(False)
        dec_mask = vals.str.fullmatch(r"[+-]?\d+\.\d+").fillna(False)
        if pct_mask.mean() >= 0.6:
            num = pd.to_numeric(vals.str.rstrip("%"), errors="coerce")
            disp[col] = num
            fmt = "%+.2f%%" if (num.dropna() < 0).any() else "%.2f%%"
            cfg[col] = st.column_config.NumberColumn(format=fmt)
        elif dec_mask.mean() >= 0.6:
            disp[col] = pd.to_numeric(vals, errors="coerce")
            cfg[col] = st.column_config.NumberColumn(format="%.2f")

    if title:
        st.markdown(f"**{title}**（{len(disp)} 条）")
    h = height if height is not None else min(420, 35 * len(disp) + 40)
    st.dataframe(disp, use_container_width=True, hide_index=True,
                 height=h, column_config=cfg)


def show_next_steps(links: List[tuple]):
    """渲染「下一步」引导卡片。

    Args:
        links: [(page_path, label, icon), ...]
    """
    if not links:
        return
    st.markdown("##### 下一步")
    cols = st.columns(len(links))
    for i, (page, label, icon) in enumerate(links):
        with cols[i]:
            st.page_link(page, label=label, icon=icon, use_container_width=True)


def show_empty_state(message: str, icon: str = "👆"):
    """统一的空状态/操作引导提示。"""
    st.info(f"{icon} {message}")


# ============================================================
# 盘中估值组件
# ============================================================

_MODE_LABELS = {
    "holdings_penetration": "重仓穿透",
    "etf_linkage": "ETF联接追踪",
    "official_fallback": "官方估算",
}

_CONF_LABELS = {
    "high": "🟢 高",
    "medium": "🟡 中",
    "low": "🔴 低",
}


def render_intraday_nav(fund_code: str, *, expanded: bool = True, show_table: bool = True):
    """
    渲染单只基金的盘中估值卡片（重仓穿透 / ETF联接 / 官方回退）。
    """
    from core.intraday_nav import get_intraday_engine
    from .utils import safe_float as _sf

    code = str(fund_code).strip().zfill(6)
    if len(code) != 6:
        st.warning("请输入 6 位基金代码")
        return None

    result = get_intraday_engine().estimate(code)
    d = result.to_dict()

    pct = d["estimated_pct"]
    color = "#DC2626" if pct >= 0 else "#16A34A"
    mode_label = _MODE_LABELS.get(d["mode"], d["mode"])
    conf = _CONF_LABELS.get(d["confidence"], d["confidence"])

    with st.expander(f"📡 盘中估值 · {d.get('fund_name') or code}", expanded=expanded):
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("估算净值", f"{d['estimated_nav']:.4f}" if d["estimated_nav"] else "—",
                  delta=f"{pct:+.2f}%" if pct else None)
        c2.metric("基准净值", f"{d['base_nav']:.4f}" if d["base_nav"] else "—")
        c3.metric("估值模式", mode_label)
        c4.metric("置信度", conf)

        if d.get("underlying_etf"):
            st.caption(f"🔗 底层标的: {d['underlying_etf']}")
        if d.get("note"):
            st.caption(d["note"])
        st.caption(f"🕐 {d['updated_at']}")

        official = d.get("official_estimate") or {}
        if official and d["mode"] != "official_fallback":
            off_pct = official.get("估算涨幅%")
            if off_pct is not None:
                st.caption(f"对照官方估算: {_sf(off_pct):+.2f}% · 净值 {_sf(official.get('估算净值', 0)):.4f}")

        if show_table and d.get("contributions"):
            import pandas as pd
            df = pd.DataFrame(d["contributions"])
            if not df.empty:
                st.markdown("**穿透明细**")
                show_sortable_df(df.rename(columns={
                    "code": "代码",
                    "name": "名称",
                    "weight_pct": "权重%",
                    "change_pct": "涨跌%",
                    "contribution_pct": "贡献%",
                }))

    return result


# ============================================================
# 按钮
# ============================================================

def render_analysis_button(label: str = "🚀 开始 AI 分析", disabled: bool = False,
                           key: Optional[str] = None) -> bool:
    """渲染分析触发按钮。"""
    return st.button(label, type="primary", disabled=disabled, use_container_width=True, key=key)


# ============================================================
# 表单控件工厂
# ============================================================

def render_param_form(parameters: List[Dict[str, Any]]) -> Dict[str, Any]:
    """根据模板参数动态生成表单控件。"""
    values = {}
    for param in parameters:
        name = param["name"]
        label = param.get("label", name)
        ptype = param.get("type", "text")
        required = param.get("required", False)
        default = param.get("default", "")
        label_text = f"{label}{' *' if required else ''}"
        key = f"param_{name}"

        if ptype == "text":
            values[name] = st.text_area(label_text, value=default, height=120,
                                        key=key, help=param.get("description", ""))
        elif ptype == "string":
            values[name] = st.text_input(label_text, value=default,
                                         key=key, help=param.get("description", ""))
        elif ptype == "select":
            options = param.get("options", [])
            values[name] = st.selectbox(label_text, options=options,
                                        key=key, help=param.get("description", ""))
        else:
            values[name] = st.text_input(label_text, value=str(default),
                                         key=key, help=param.get("description", ""))
    return values


# ============================================================
# 顶部状态条：数据刷新 / 模型 / 版本
# ============================================================

def render_top_toolbar(show_ai_config: bool = True):
    """
    在主内容区顶部渲染数据刷新状态 + 刷新按钮。

    模型/版本信息统一在侧边栏底部展示，此处只负责“数据新鲜度 + 手动刷新”，
    避免与侧边栏、页面头部重复。

    Args:
        show_ai_config: 保留参数以兼容既有调用，当前不再影响渲染。
    """
    import time as _t

    ts_key = "global_data_refresh_ts"
    if ts_key not in st.session_state:
        st.session_state[ts_key] = _t.time()
    elapsed = _t.time() - st.session_state[ts_key]
    if elapsed < 60: ago = "刚刚"
    elif elapsed < 3600: ago = f"{int(elapsed//60)}分钟前"
    else: ago = f"{int(elapsed//3600)}小时前"

    # 紧凑单行：数据新鲜度与刷新按钮相邻靠左，右侧留白填充
    c_text, c_btn, _spacer = st.columns([1.5, 0.7, 5], vertical_alignment="center")
    with c_text:
        st.caption(f"🕐 数据刷新: {ago}")
    with c_btn:
        if st.button("🔄 刷新", key="btn_global_refresh", help="清除全部数据缓存重新获取"):
            st.session_state[ts_key] = _t.time()
            st.cache_data.clear()
            st.rerun()

    st.divider()


# ============================================================
# 侧边栏：全局配置
# ============================================================

def render_sidebar_config(show_ai_config: bool = True) -> Dict[str, Any]:
    """
    在各分析页面侧边栏渲染全局搜索 + AI 模型与深度配置。

    Args:
        show_ai_config: 是否显示 AI 模型选择器（非 AI 页面可传 False）
    """
    config_values = {}

    with st.sidebar:
        # ---- 全局基金速查 ----
        with st.expander("🔍 基金速查", expanded=True):
            search_code = st.text_input(
                "输入基金代码或名称",
                placeholder="000001 / 华夏成长",
                key="global_fund_search",
                label_visibility="collapsed",
            )
            if search_code and len(search_code.strip()) >= 2:
                try:
                    from core.data_fetcher import get_fetcher
                    fetcher = get_fetcher()
                    results = fetcher.search_fund(search_code.strip())
                    if results is not None and not results.empty:
                        for _, row in results.head(5).iterrows():
                            code = str(row.get("基金代码", ""))
                            name = str(row.get("基金简称", row.get("基金名称", "")))
                            st.caption(f"📌 `{code}` {name[:12]}")
                    else:
                        st.caption("未找到匹配基金")
                except Exception:
                    pass

        # ---- AI 引擎 ----
        if show_ai_config:
            st.markdown("**AI 引擎**")
            from core.config import config as app_config

            models = app_config.available_models
            help_text = (
                f"{models[0]}: 快速分析 | {models[-1]}: 深度推理"
                if len(models) >= 2 else str(models)
            )

            # 从云端偏好恢复默认模型/深度
            pref_model_idx = 0
            pref_depth = "标准"
            try:
                if st.session_state.get("user"):
                    from core.user_workspace import get_prefs
                    prefs = get_prefs().get("prefs") or {}
                    pm = prefs.get("default_model")
                    if pm in models:
                        pref_model_idx = models.index(pm)
                    if prefs.get("default_depth") in ("简要", "标准", "深度", "极致"):
                        pref_depth = prefs["default_depth"]
            except Exception:
                pass

            config_values["model"] = st.selectbox(
                "模型", options=models, index=pref_model_idx, help=help_text, key="sidebar_ai_model"
            )
            # 让当次选择立即作用于分析引擎（此前只写偏好、不生效）
            try:
                from core.user_integrations import set_session_model
                set_session_model(config_values["model"])
            except Exception:
                pass

            config_values["depth"] = st.select_slider(
                "分析深度", options=["简要", "标准", "深度", "极致"], value=pref_depth,
                help="影响输出详细程度和 token 消耗", key="sidebar_ai_depth",
            )
            # 写回偏好（仅变更时）
            try:
                if st.session_state.get("user"):
                    from core.user_workspace import get_prefs, save_prefs
                    cur = get_prefs().get("prefs") or {}
                    if (
                        cur.get("default_model") != config_values["model"]
                        or cur.get("default_depth") != config_values["depth"]
                    ):
                        save_prefs(prefs_patch={
                            "default_model": config_values["model"],
                            "default_depth": config_values["depth"],
                        })
            except Exception:
                pass

        return config_values


# ============================================================
# 数据新鲜度 & 智能建议
# ============================================================

import time as _time

def render_data_freshness(timestamp: float = None, label: str = "数据"):
    """显示数据新鲜度标签。"""
    if timestamp is None:
        timestamp = _time.time()
    elapsed = _time.time() - timestamp
    if elapsed < 30:
        badge = '<span style="color:#10B981;">🟢 刚刚更新</span>'
    elif elapsed < 300:
        badge = f'<span style="color:#3B82F6;">🔵 {int(elapsed//60)}分钟前</span>'
    elif elapsed < 3600:
        badge = f'<span style="color:#F59E0B;">🟡 {int(elapsed//60)}分钟前</span>'
    else:
        badge = f'<span style="color:#EF4444;">🔴 {int(elapsed//3600)}小时前</span>'
    st.caption(f"{badge} · {label}", unsafe_allow_html=True)


def render_filter_suggestion(stats: dict, current_params: dict) -> str:
    """
    筛选无结果时生成智能放宽建议。

    Args:
        stats: {"f1": 淘汰数, ...} 过滤统计
        current_params: {"min_sharpe": 0.8, "max_drawdown": -30, ...} 当前参数

    Returns:
        建议文本
    """
    suggestions = []
    param_map = {
        "f1": ("夏普比率", "min_sharpe", 0.1, lambda v: f"降至 {v-0.1:.1f}"),
        "f3": ("最大回撤", "max_drawdown", -5, lambda v: f"放宽至 {v+5}%"),
        "f5": ("索提诺比率", "min_sortino", 0.2, lambda v: f"降至 {v-0.2:.1f}"),
        "f7": ("超额收益", "min_excess", -1, lambda v: f"降至 {v-1}%"),
        "f8": ("规模限制", "", 0, lambda v: "暂时关闭规模过滤"),
    }

    # 按淘汰数排序，找到阻碍最大的过滤
    filter_stats = {k: v for k, v in stats.items() if isinstance(k, str) and k.startswith("f")}
    sorted_filters = sorted(filter_stats.items(), key=lambda x: x[1], reverse=True)

    for fkey, count in sorted_filters[:3]:
        if count > 5 and fkey in param_map:
            label, pkey, delta, sug = param_map[fkey]
            if pkey in current_params:
                new_val = current_params[pkey] + delta
                suggestions.append(f"→ 将**{label}**{sug(new_val)}（可多通过 ~{count} 只）")

    if suggestions:
        return "### 💡 一键优化建议\n\n" + "\n".join(suggestions)
    return ""


# ============================================================
# 结果展示
# ============================================================

def show_token_usage(usage: Dict[str, int]):
    """展示 Token 用量（紧凑单行 caption，避免三张大卡占满整行）。"""
    if usage and usage.get("total_tokens", 0) > 0:
        st.caption(
            f"📊 Token · 输入 {usage['input_tokens']:,} / "
            f"输出 {usage['output_tokens']:,} / 总计 {usage['total_tokens']:,}"
        )


def show_elapsed_time(seconds: float):
    """展示耗时。"""
    if seconds > 0:
        st.caption(f"⏱️ {seconds:.1f}s")


def show_run_meta(elapsed_seconds: float = 0, usage: Optional[Dict[str, int]] = None):
    """一行展示本次调用的耗时 + Token 用量（合并 show_elapsed_time + show_token_usage）。"""
    parts = []
    if elapsed_seconds and elapsed_seconds > 0:
        parts.append(f"⏱️ {elapsed_seconds:.1f}s")
    if usage and usage.get("total_tokens", 0) > 0:
        parts.append(
            f"📊 Token 输入 {usage['input_tokens']:,} / "
            f"输出 {usage['output_tokens']:,} / 总计 {usage['total_tokens']:,}"
        )
    if parts:
        st.caption("　·　".join(parts))


def rating_badge(value: float, metric: str) -> str:
    """自动评级徽章。根据指标值返回HTML徽章。"""
    thresholds = {
        "sharpe": [(2.0, "🏅 卓越", "#7C3AED"), (1.2, "🥈 优秀", "#3B82F6"), (0.7, "🥉 良好", "#10B981"), (-99, "⚪ 一般", "#94A3B8")],
        "sortino": [(2.5, "🏅", "#7C3AED"), (1.5, "🥈", "#3B82F6"), (0.8, "🥉", "#10B981"), (-99, "⚪", "#94A3B8")],
        "max_dd": [(-5, "🛡️ 极低回撤", "#10B981"), (-15, "✅ 可控", "#3B82F6"), (-25, "⚠️ 中等", "#F59E0B"), (-99, "🔴 高回撤", "#DC2626")],
        "return": [(30, "🔥 顶级", "#DC2626"), (15, "📈 优秀", "#F59E0B"), (5, "📊 良好", "#3B82F6"), (-99, "📉", "#94A3B8")],
        "win_rate": [(70, "🏅", "#7C3AED"), (55, "🥈", "#3B82F6"), (45, "🥉", "#10B981"), (-99, "⚪", "#94A3B8")],
    }
    rules = thresholds.get(metric, [(999, "", "#94A3B8")])
    for threshold, label, color in rules:
        if value >= threshold:
            return f'<span style="display:inline-block;padding:2px 8px;border-radius:12px;background:{color}20;color:{color};font-size:0.75rem;font-weight:600;margin-left:6px;">{label}</span>'
    return ""


def show_error(message: str):
    st.error(message)


def show_warning_box(message: str):
    st.warning(message)


def show_info_box(message: str):
    st.info(message)


# ============================================================
# 安全执行边界
# ============================================================


def safe_run(func, error_prefix: str = "操作"):
    """
    安全执行包装器 — 捕获所有异常，显示友好错误而非白屏。

    用法:
        safe_run(lambda: fetcher.get_fund_list("股票型"), "获取基金列表")
    """
    import traceback
    try:
        return func()
    except Exception as e:
        tb = traceback.format_exc()
        st.error(f"❌ {error_prefix}失败: {str(e)[:300]}")
        with st.expander("🔍 错误详情（调试用）", expanded=False):
            st.code(tb[-2000:], language="text")
        return None


def run_analysis_stream(analyzer, system_prompt, user_prompt,
                        max_tokens: int = 4096, temperature: float = 0.3,
                        placeholder_label: str = "📝 分析报告",
                        page_name: str = "", fund_code: str = "", fund_name: str = "",
                        enable_decision_dashboard: bool = True):
    """
    统一的流式分析执行器 — 消除10个页面中的重复 pattern。

    用法:
        result = run_analysis_stream(analyzer, sys_prompt, usr_prompt, max_tokens=4096)
        if result: show_token_usage(result["usage"])

    可选 page_name/fund_code/fund_name 用于自动保存到分析历史。
    enable_decision_dashboard: 解析并渲染结构化决策仪表盘 JSON。

    配额：本函数是所有流式 AI 分析的唯一出口，因此套餐限额在这里统一拦截。
          社区版（EDITION=community）不做任何限制。
    """
    # ---- 套餐配额守卫（托管版生效，社区版直通）----
    try:
        from core.middleware import require_ai_quota
        if not require_ai_quota(placeholder_label or "AI 分析"):
            return {
                "success": False,
                "content": "",
                "model": "",
                "usage": {},
                "elapsed_seconds": 0,
                "error": "今日 AI 分析额度已用完",
                "quota_exceeded": True,
            }
    except ImportError:
        pass

    from core.decision_extractor import DECISION_JSON_INSTRUCTION, extract_dashboard, merge_with_meta
    from core.decision_renderer import render_decision_dashboard

    if enable_decision_dashboard and "schema_version" not in system_prompt:
        system_prompt = system_prompt.rstrip() + "\n" + DECISION_JSON_INSTRUCTION

    import html as _html
    import time as _time

    st.markdown(f"#### {placeholder_label}")
    # B：首 token 前的状态骨架 + A：推理模型思考过程展示
    status = st.status("⏳ 正在连接模型，准备生成…", expanded=False)
    reasoning_box = status.empty()
    placeholder = st.empty()
    dashboard_slot = st.empty()
    full_text = ""
    reasoning_text = ""
    final_result = None
    dashboard_data = None
    first_content = False
    _start = _time.time()

    def _render_reasoning(txt: str):
        safe = _html.escape(txt[-1500:])
        sc = surface_colors()
        reasoning_box.markdown(
            f"<div style='color:{sc['reasoning']};font-size:0.82rem;line-height:1.5;"
            "white-space:pre-wrap;max-height:220px;overflow:auto;'>"
            f"{safe}</div>",
            unsafe_allow_html=True,
        )

    try:
        stream_gen = analyzer.analyze_stream(
            system_prompt=system_prompt,
            user_prompt=user_prompt,
            max_tokens=max_tokens,
            temperature=temperature,
            meter_source=page_name or "analysis",
        )
        while True:
            try:
                chunk = next(stream_gen)
            except StopIteration as e:
                # 生成器 return 值含 token/success；始终优先采用（避免被中间 dict chunk 占位）
                if isinstance(getattr(e, "value", None), dict):
                    final_result = e.value
                break

            # 思考内容（推理模型）—— 单独灰字展示，不并入正文
            if isinstance(chunk, dict) and "reasoning" in chunk:
                if not reasoning_text:
                    status.update(label="🧠 模型思考中…", expanded=True)
                reasoning_text += chunk["reasoning"]
                _render_reasoning(reasoning_text)
                continue

            if isinstance(chunk, str):
                if not first_content:
                    first_content = True
                    status.update(label="✍️ 正在生成分析…", expanded=False)
                full_text += chunk
                placeholder.markdown(full_text + "▌")
                # 实时元信息：已生成字数
                status.update(label=f"✍️ 正在生成分析… 已 {len(full_text)} 字")
            elif isinstance(chunk, dict) and "success" in chunk:
                final_result = chunk

        placeholder.markdown(full_text)
        _elapsed = final_result.get("elapsed_seconds") if final_result else round(_time.time() - _start, 1)
        status.update(
            label=f"✅ 分析完成 · {len(full_text)} 字" + (f" · 用时 {_elapsed}s" if _elapsed else ""),
            state="complete",
            expanded=False,
        )
    except Exception as e:
        placeholder.markdown(full_text + f"\n\n> ⚠️ 分析中断: {str(e)[:200]}")
        status.update(label="⚠️ 分析中断", state="error", expanded=False)

    # ---- 先落库，再渲染仪表盘（避免仪表盘异常导致历史未保存）----
    _saved_hid = _auto_save_analysis_history(
        full_text=full_text,
        page_name=page_name,
        fund_code=fund_code,
        fund_name=fund_name,
        system_prompt=system_prompt,
        user_prompt=user_prompt,
        placeholder_label=placeholder_label,
        final_result=final_result,
        structured_result=None,
    )

    if enable_decision_dashboard and full_text:
        try:
            raw_dash = extract_dashboard(full_text)
            if raw_dash:
                model_name = ""
                if final_result:
                    model_name = final_result.get("model", "")
                dashboard_data = merge_with_meta(
                    raw_dash,
                    page_name=page_name,
                    fund_code=fund_code,
                    fund_name=fund_name,
                    model=model_name,
                )
                with dashboard_slot.container():
                    render_decision_dashboard(dashboard_data, expanded=True)
                if final_result is not None:
                    final_result["dashboard"] = dashboard_data
                try:
                    _maybe_create_signal_from_dashboard(dashboard_data, fund_code, fund_name)
                except Exception:
                    pass
                # 若已有历史 ID，补写结构化结果（失败不影响正文已保存）
                if _saved_hid and dashboard_data:
                    try:
                        from core.db import get_connection, DB_MODE
                        ph = "%s" if DB_MODE == "postgresql" else "?"
                        import json as _json
                        with get_connection() as conn:
                            conn.execute(
                                f"UPDATE analysis_history SET structured_result_json = {ph} WHERE id = {ph}",
                                (_json.dumps(dashboard_data, ensure_ascii=False), _saved_hid),
                            )
                    except Exception:
                        pass
        except Exception as _dash_err:
            st.caption(f"⚠️ 决策仪表盘渲染跳过：{str(_dash_err)[:80]}")

    # 回传正文，供页面级兜底保存
    if final_result is None:
        final_result = {"success": bool(full_text.strip()), "elapsed_seconds": 0, "usage": {}, "model": ""}
    final_result["content"] = full_text
    final_result["history_id"] = _saved_hid

    # ---- 分析结果操作栏 ----
    if full_text.strip() and (final_result.get("success") or full_text.strip()):
        import hashlib as _hlib
        import re as _re
        _key_seed = _hlib.md5(full_text[:200].encode()).hexdigest()[:12]

        # 构建有意义的文件名：页面上报 → 基金信息 → placeholder标签
        _name_parts = []
        if page_name:
            _name_parts.append(page_name)
        if fund_name:
            _name_parts.append(fund_name)
        elif fund_code:
            _name_parts.append(fund_code)
        if not _name_parts:
            # 从 placeholder_label 提取（去掉 emoji）
            _label = _re.sub(r'[^\w一-鿿]', '', placeholder_label.split()[-1] if placeholder_label else "")
            _name_parts.append(_label or "分析报告")
        _base_name = "_".join(_name_parts)
        _base_name = _re.sub(r'[<>:"/\\|?*]', '_', _base_name)[:80]

        st.markdown("---")
        action_cols = st.columns([1, 1, 1, 2])
        with action_cols[0]:
            st.download_button(
                label="📄 导出 Word",
                data=_export_result_as(full_text, "word"),
                file_name=f"{_base_name}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key=f"export_word_{_key_seed}",
            )
        with action_cols[1]:
            st.download_button(
                label="📊 导出 Excel",
                data=_export_result_as(full_text, "excel"),
                file_name=f"{_base_name}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key=f"export_xlsx_{_key_seed}",
            )
        with action_cols[2]:
            save_key = f"save_hist_{_key_seed}"
            if _saved_hid is not None:
                st.caption(f"📚 已自动保存 #{_saved_hid}")
            if st.button("📚 再存一份", key=save_key):
                try:
                    from core.history import get_history_manager
                    hm = get_history_manager()
                    hid = hm.save(
                        page_name=page_name or placeholder_label or "未知页面",
                        fund_code=fund_code or "",
                        fund_name=fund_name or "",
                        system_prompt=system_prompt,
                        user_prompt=user_prompt,
                        result_content=full_text,
                        usage=final_result.get("usage", {}),
                        elapsed_seconds=final_result.get("elapsed_seconds", 0),
                        model=final_result.get("model", ""),
                        structured_result=final_result.get("dashboard"),
                    )
                    st.toast(f"✅ 已再次保存 (#{hid})", icon="📚")
                except Exception as e:
                    st.toast(f"⚠️ 保存失败：{str(e)[:80]}", icon="⚠️")
        with action_cols[3]:
            st.page_link("pages/24_导出中心.py", label="📤 导出中心", icon="📤")

    # 记录 API 调用用量：正文已生成即计次
    if full_text.strip():
        try:
            from core.middleware import record_api_call, QUOTA_MAP
            if not record_api_call():
                st.caption("⚠️ 用量计量未写入（未登录或数据库异常）")
            else:
                user = st.session_state.get("user") or {}
                if user.get("id"):
                    from core.user_repo import UserRepo
                    usage = UserRepo().get_usage(int(user["id"]))
                    tier = usage.get("tier") or user.get("tier") or "free"
                    used = int(usage.get("api_calls_today", 0) or 0)
                    limit = int(QUOTA_MAP.get(tier, {}).get("ai_analysis") or usage.get("api_calls_limit") or 0)
                    if limit and used / limit >= 0.8:
                        st.warning(f"⚠️ AI 用量已达今日配额的 {used}/{limit}，请注意节省。")
                    elif limit:
                        st.caption(f"⚡ 今日剩余 AI 次数约 {max(0, limit - used)}")
        except Exception:
            st.caption("⚠️ 用量计量异常")

    return final_result


def _auto_save_analysis_history(
    *,
    full_text: str,
    page_name: str,
    fund_code: str,
    fund_name: str,
    system_prompt: str,
    user_prompt: str,
    placeholder_label: str,
    final_result: dict = None,
    structured_result: dict = None,
):
    """正文一旦生成即写入分析历史，返回 history_id 或 None。"""
    if not (full_text or "").strip():
        return None
    _auto_ok = True
    try:
        if st.session_state.get("user"):
            from core.user_workspace import get_prefs
            _auto_ok = bool((get_prefs().get("prefs") or {}).get("auto_save_history", True))
    except Exception:
        _auto_ok = True
    if not _auto_ok:
        st.caption("ℹ️ 已关闭「自动保存历史」，可在账户偏好中开启")
        return None

    _hist_page = (page_name or placeholder_label or "AI分析").strip() or "AI分析"
    # 用稳定 md5，避免依赖 hash() 随机化；同文多次点击不重复塞 session 防抖
    import hashlib as _hl
    _auto_hist_key = "auto_hist_" + _hl.md5(
        f"{_hist_page}|{fund_code}|{full_text[:240]}".encode("utf-8", "ignore")
    ).hexdigest()[:16]
    if _auto_hist_key in st.session_state and st.session_state[_auto_hist_key] is not None:
        hid = st.session_state[_auto_hist_key]
        st.caption(f"📚 已自动保存到分析历史 (#{hid})")
        return hid
    try:
        from core.history import get_history_manager
        from core.user_workspace import touch_recent
        _usage = (final_result or {}).get("usage", {}) or {}
        # 超长 prompt 截断，避免异常拖垮保存
        _sp = (system_prompt or "")[:4000]
        _up = (user_prompt or "")[:8000]
        hid = get_history_manager().save(
            page_name=_hist_page,
            fund_code=fund_code or "",
            fund_name=fund_name or "",
            system_prompt=_sp,
            user_prompt=_up,
            result_content=full_text,
            usage=_usage,
            elapsed_seconds=(final_result or {}).get("elapsed_seconds", 0) or 0,
            model=(final_result or {}).get("model", "") or "",
            structured_result=structured_result,
        )
        st.session_state[_auto_hist_key] = hid
        st.success(f"📚 已自动保存到分析历史 (#{hid})")
        if fund_code:
            try:
                touch_recent("fund", fund_code, fund_name or fund_code)
            except Exception:
                pass
        return hid
    except Exception as e:
        st.session_state[_auto_hist_key] = None
        st.warning(f"⚠️ 自动保存历史失败：{type(e).__name__}: {str(e)[:160]}")
        return None


def _maybe_create_signal_from_dashboard(dashboard: dict, fund_code: str, fund_name: str):
    """从 dashboard 的 signal_for_validation 自动写入信号池（仅当有基金代码时）。"""
    if not fund_code or not dashboard:
        return
    sfv = dashboard.get("signal_for_validation") or {}
    action = sfv.get("action")
    if not action:
        return
    try:
        from core.signals import get_signal_manager
        sm = get_signal_manager()
        sm.create(
            fund_code=fund_code,
            fund_name=fund_name or "",
            action=action,
            confidence=float(sfv.get("confidence", 0.5)),
            score=float(sfv.get("confidence", 0.5)) * 100,
            reason=dashboard.get("core_conclusion", {}).get("one_liner", "来自决策仪表盘"),
        )
    except Exception:
        pass


def _export_result_as(text: str, fmt: str) -> bytes:
    """快速导出分析结果为 Word 或 Excel。"""
    import io
    if fmt == "word":
        try:
            from docx import Document
        except ImportError as e:
            raise ImportError("需要安装 word 导出依赖: pip install python-docx") from e
        doc = Document()
        doc.add_heading("基金分析报告", level=0)
        for line in text.split("\n"):
            line = line.strip()
            if line:
                doc.add_paragraph(line)
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
    if fmt == "excel":
        try:
            from openpyxl import Workbook
        except ImportError as e:
            raise ImportError("需要安装 Excel 导出依赖: pip install openpyxl") from e
        wb = Workbook()
        ws = wb.active
        ws.title = "分析报告"
        ws.cell(row=1, column=1, value="基金分析报告")
        for i, line in enumerate(text.split("\n"), 2):
            ws.cell(row=i, column=1, value=line.strip())
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    return text.encode("utf-8")


# ============================================================
# 页面初始化
# ============================================================

def init_page_state(defaults: Dict[str, Any] = None):
    if defaults:
        for key, val in defaults.items():
            if key not in st.session_state:
                st.session_state[key] = val


def check_api_ready() -> bool:
    """
    检查生效的 LLM 配置是否可用。

    未配置时给出**两条**解决路径：界面自助配置（推荐给终端用户）
    与服务端 .env 配置（推荐给自部署者）。本地推理服务不需要 API Key。
    """
    from core.config import config as app_config
    from core.llm_presets import get_preset
    from core.user_integrations import resolve_llm_config

    if app_config.is_configured:
        return True

    eff = resolve_llm_config()
    preset = get_preset(eff.preset)
    label = preset.label if preset else eff.preset

    st.error(f"⚠️ **{label}** 尚未配置完成，AI 分析功能暂不可用。")
    st.markdown(
        "**方式一（推荐）**：在页面左侧进入 **账户设置 → 🔌 AI 与数据源**，"
        "填写自己的 API Key 并点「测试连通性」，保存后立即生效，无需重启。"
    )
    if preset and preset.api_key_url:
        st.caption(f"🔑 获取密钥：{preset.api_key_url}")
    if preset and not preset.requires_key:
        st.caption(f"ℹ️ 当前服务商为本地推理服务，通常无需 API Key，请确认已填写模型名称与接口地址。")

    st.markdown(
        "**方式二（自部署）**：在服务端 `.env` 中设置 `LLM_PRESET` / `LLM_API_KEY` / "
        "`LLM_MODEL` 三项后重启服务。"
    )
    st.caption(
        "💡 未配置 AI 也不影响使用：基金筛选、定投回测、费率计算、蒙特卡洛等"
        "纯本地计算功能完全可用。"
    )
    return False
